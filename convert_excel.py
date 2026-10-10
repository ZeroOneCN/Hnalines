#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
海航 PLUS 适用航班 Excel → 网站 JSON 数据转换工具

用法:
    python convert_excel.py [Excel文件路径] [-o 输出文件路径]

默认读取当前目录下最新的 *.xlsx（排除临时文件），生成 hainan_airlines_data.json。
适用于任一航季的官方 Excel（工作表结构：出港城市/进港城市/航班号/班期/出发时刻/适用日期段/适用产品）。
"""

import json
import glob
import os
import re
import sys
from openpyxl import load_workbook

DEFAULT_OUTPUT = 'hainan_airlines_data.json'

# 城市 → 省份映射（含机场/别名；转换新航季数据时若出现新城市，在此补充）
PROVINCE_MAPPING = {
    '北京': '北京', '上海': '上海', '天津': '天津', '重庆': '重庆',
    '乌鲁木齐': '新疆', '阿克苏': '新疆', '喀什': '新疆', '哈密': '新疆', '塔城': '新疆', '那拉提': '新疆',
    '图木舒克': '新疆', '博乐': '新疆', '库尔勒': '新疆', '库车': '新疆', '伊宁': '新疆', '于田': '新疆',
    '和田': '新疆', '石河子': '新疆', '阿勒泰': '新疆', '阿拉尔': '新疆', '楼兰': '新疆',
    '西安': '陕西', '咸阳': '陕西', '延安': '陕西', '榆林': '陕西', '汉中': '陕西', '安康': '陕西',
    '郑州': '河南', '洛阳': '河南', '安阳': '河南', '南阳': '河南', '信阳': '河南',
    '合肥': '安徽', '安庆': '安徽', '黄山': '安徽', '池州': '安徽', '亳州': '安徽', '阜阳': '安徽', '蚌埠': '安徽',
    '海口': '海南', '三亚': '海南', '琼海': '海南', '三沙': '海南',
    '杭州': '浙江', '宁波': '浙江', '温州': '浙江', '台州': '浙江', '舟山': '浙江', '丽水': '浙江', '义乌': '浙江', '衢州': '浙江',
    '成都': '四川', '天府': '四川', '绵阳': '四川', '泸州': '四川', '巴中': '四川', '南充': '四川', '达州': '四川',
    '宜宾': '四川', '广元': '四川', '西昌': '四川', '阆中': '四川',
    '大理': '云南', '西双版纳': '云南', '昆明': '云南', '芒市': '云南', '丽江': '云南', '腾冲': '云南',
    '香格里拉': '云南', '沧源': '云南', '澜沧': '云南', '保山': '云南', '文山': '云南', '昭通': '云南',
    '临沧': '云南', '德宏': '云南', '普洱': '云南', '迪庆': '云南',
    '大连': '辽宁', '沈阳': '辽宁', '锦州': '辽宁', '朝阳': '辽宁', '丹东': '辽宁', '营口': '辽宁',
    '青岛': '山东', '济南': '山东', '临沂': '山东', '威海': '山东', '烟台': '山东', '潍坊': '山东',
    '日照': '山东', '东营': '山东', '济宁': '山东', '菏泽': '山东',
    '广州': '广东', '深圳': '广东', '惠州': '广东', '珠海': '广东', '揭阳': '广东', '湛江': '广东', '韶关': '广东', '佛山': '广东', '梅州': '广东',
    '南京': '江苏', '苏州': '江苏', '常州': '江苏', '无锡': '江苏', '南通': '江苏', '徐州': '江苏',
    '扬州': '江苏', '淮安': '江苏', '盐城': '江苏', '连云港': '江苏', '无锡硕放': '江苏', '扬州泰州': '江苏',
    '长沙': '湖南', '张家界': '湖南', '衡阳': '湖南', '湘西': '湖南', '郴州': '湖南', '邵阳': '湖南', '永州': '湖南', '岳阳': '湖南', '怀化': '湖南', '常德': '湖南',
    '哈尔滨': '黑龙江', '鸡西': '黑龙江', '大庆': '黑龙江', '齐齐哈尔': '黑龙江', '牡丹江': '黑龙江', '佳木斯': '黑龙江',
    '长春': '吉林', '吉林': '吉林', '延吉': '吉林', '通化': '吉林', '白山': '吉林',
    '石家庄': '河北', '秦皇岛': '河北', '唐山': '河北', '邯郸': '河北', '邢台': '河北', '张家口': '河北', '承德': '河北',
    '太原': '山西', '大同': '山西', '长治': '山西', '吕梁': '山西', '运城': '山西', '临汾': '山西', '晋城': '山西',
    '呼和浩特': '内蒙古', '包头': '内蒙古', '二连浩特': '内蒙古', '阿尔山': '内蒙古', '满洲里': '内蒙古', '呼伦贝尔': '内蒙古',
    '霍林郭勒': '内蒙古', '乌海': '内蒙古', '乌兰浩特': '内蒙古', '通辽': '内蒙古', '赤峰': '内蒙古', '鄂尔多斯': '内蒙古',
    '锡林浩特': '内蒙古', '扎兰屯': '内蒙古', '乌兰察布': '内蒙古', '巴彦淖尔': '内蒙古',
    '兰州': '甘肃', '敦煌': '甘肃', '张掖': '甘肃', '金昌': '甘肃', '庆阳': '甘肃', '陇南': '甘肃', '嘉峪关': '甘肃', '天水': '甘肃',
    '西宁': '青海', '玉树': '青海', '格尔木': '青海', '德令哈': '青海',
    '银川': '宁夏', '中卫': '宁夏', '固原': '宁夏',
    '拉萨': '西藏', '昌都': '西藏', '林芝': '西藏', '日喀则': '西藏', '阿里': '西藏',
    '南宁': '广西', '桂林': '广西', '北海': '广西', '柳州': '广西', '玉林': '广西', '百色': '广西', '梧州': '广西',
    '贵阳': '贵州', '茅台': '贵州', '遵义': '贵州', '铜仁': '贵州', '兴义': '贵州', '安顺': '贵州', '毕节': '贵州', '凯里': '贵州', '六盘水': '贵州', '荔波': '贵州',
    '遵义新舟': '贵州', '遵义茅台': '贵州', '新舟': '贵州',
    '福州': '福建', '厦门': '福建', '泉州': '福建', '武夷山': '福建', '三明': '福建', '龙岩': '福建', '宁德': '福建',
    '南昌': '江西', '赣州': '江西', '九江': '江西', '宜春': '江西', '上饶': '江西', '井冈山': '江西', '景德镇': '江西', '吉安': '江西',
    '武汉': '湖北', '宜昌': '湖北', '荆州': '湖北', '襄阳': '湖北', '恩施': '湖北', '十堰': '湖北', '神农架': '湖北',
    '北京首都': '北京', '北京大兴': '北京', '大兴': '北京', '首都': '北京',
    '上海浦东': '上海', '上海虹桥': '上海', '浦东': '上海', '虹桥': '上海',
    '成都天府': '四川', '成都双流': '四川', '重庆江北': '重庆', '江北': '重庆', '万州': '重庆',
}


def get_province(city: str) -> str:
    """根据城市名获取所属省份。

    参数: city - 城市或机场名称
    返回值: 省份字符串，未匹配返回 '其他'
    """
    for key, prov in PROVINCE_MAPPING.items():
        if key in city:
            return prov
    return '其他'


def convert_product(raw: str) -> str:
    """将 Excel「适用产品」描述转换为站点产品标记。

    规则: '适用2666' → '2666'；'适用666或2666' → '666/2666'。
    注意必须先判断是否含独立 666 权益（'或'字样），避免 '2666' 子串误判为含 666。

    参数: raw - Excel 原始产品描述
    返回值: 产品标记字符串（如 '2666' 或 '666/2666'）
    """
    text = raw.replace('适用', '').replace('或', '/')
    if text == '2666':
        return '2666'
    if text == '666/2666':
        return '666/2666'
    tokens = []
    if re.search(r'(?<!2)666', text):
        tokens.append('666')
    if '2666' in text:
        tokens.append('2666')
    return '/'.join(tokens) if tokens else text


def normalize_time(t: str) -> str:
    """规范出发时刻为 HH:MM 格式。

    参数: t - 原始时刻字符串（如 '7:45'）
    返回值: 规范后的时刻字符串；无法解析时原样返回
    """
    m = re.match(r'^(\d{1,2}):(\d{2})$', str(t).strip())
    if not m:
        return str(t).strip()
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def find_source_xlsx(path: str = None) -> str:
    """定位待转换的 Excel 源文件。

    参数: path - 命令行指定的文件路径；为空时自动取目录下最新 *.xlsx
    返回值: Excel 文件路径
    """
    if path:
        return path
    candidates = [f for f in glob.glob('*.xlsx') if not f.startswith('~$')]
    if not candidates:
        raise FileNotFoundError('未找到 Excel 源文件，请指定路径')
    return max(candidates, key=os.path.getmtime)


def convert(src: str, out: str, data_version: str = None, season: str = None) -> int:
    """执行转换：读取全部工作表 → 按行去重 → 生成站点 JSON。

    保留 Excel 原始行序（官方已按出港城市拼音 A-Z 排序），不做重排序。

    参数: src - Excel 源路径; out - JSON 输出路径
          data_version - 数据版本标识（默认从文件名推断，如 '2026冬春航季'）
          season - 航季适用日期描述
    返回值: 转换记录条数
    """
    wb = load_workbook(src, read_only=True, data_only=True)
    flights = []
    seen = set()

    def add_row(r):
        """处理单行记录：清洗字段、去重并追加到结果列表。"""
        vals = [str(c).strip() if c is not None else '' for c in list(r)[:7]]
        vals += [''] * (7 - len(vals))
        dep, arr, fno, sched, time, dates, prod = vals
        if not fno or fno == '航班号' or not dep or not arr:
            return
        key = f"{fno}|{dep}|{arr}|{sched}|{time}|{dates}|{prod}"
        if key in seen:
            return
        seen.add(key)
        flights.append({
            'flight_number': fno,
            'departure_city': dep,
            'arrival_city': arr,
            'departure_time': normalize_time(time),
            'schedule': sched,
            'date_range': dates,
            'product': convert_product(prod),
            'departure_province': get_province(dep),
            'arrival_province': get_province(arr),
        })

    for name in wb.sheetnames:
        for r in wb[name].iter_rows(values_only=True):
            add_row(r)

    if data_version is None:
        m = re.search(r'(\d{2})冬春|(\d{2})夏秋', os.path.basename(src))
        data_version = f"20{m.group(1) or m.group(2)}{'冬春' if m.group(1) else '夏秋'}航季" if m else '未知版本'
    if season is None:
        season = ''

    data = {
        'data_version': data_version,
        'season': season,
        'total_flights': len(flights),
        'flights': flights,
    }
    with open(out, 'w', encoding='utf-8') as fp:
        json.dump(data, fp, ensure_ascii=False, indent=2)

    unmapped = sorted({c for f in flights for c in (f['departure_city'], f['arrival_city'])
                       if get_province(c) == '其他'})
    print(f"源文件: {src}")
    print(f"转换记录: {len(flights)} 条（去重后） → {out}")
    products = {}
    for f in flights:
        products[f['product']] = products.get(f['product'], 0) + 1
    print(f"产品分布: {products}")
    print(f"未匹配省份城市: {unmapped if unmapped else '无'}")
    return len(flights)


def main():
    """命令行入口：解析参数并执行转换。"""
    args = [a for a in sys.argv[1:] if not a.startswith('-')]
    out = DEFAULT_OUTPUT
    if '-o' in sys.argv:
        out = sys.argv[sys.argv.index('-o') + 1]
    src = find_source_xlsx(args[0] if args else None)
    convert(src, out)


if __name__ == '__main__':
    main()
